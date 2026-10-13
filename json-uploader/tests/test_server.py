"""Pruebas del servidor. Desde json-uploader/:  python -m unittest discover tests"""

import base64
import csv
import io
import json
import logging
import os
import re
import shutil
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook

# server.py lee su configuración del entorno al importarse
UPLOADS = tempfile.mkdtemp(prefix="codea2-pruebas-")
os.environ.update(UPLOAD_DIR=UPLOADS, UPLOAD_API_KEY="clave-visor", DASHBOARD_PASSWORD="clave-panel",
                  DASHBOARD_PUBLIC="")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import server  # noqa: E402
from codea2 import model  # noqa: E402

logging.disable(logging.CRITICAL)

# Sesiones con el formato anterior al 06/10/2026: uploads/ sigue lleno de ellas
SAMPLES = Path(__file__).resolve().parent / "fixtures"
CURRENT = Path(__file__).resolve().parents[2] / "JSON Samples"
VISOR = {"X-API-Key": "clave-visor"}
BROWSER = {"Accept": "text/html"}


def sample(name):
    return (SAMPLES / name).read_bytes()


def other_start(content):
    """La misma sesión, como si la hubiera jugado otro visor a otra hora."""
    data = json.loads(content)
    data["startedUtc"] = "2026-10-01T10:00:00.0000000Z"
    return json.dumps(data).encode()


class ServerCase(unittest.TestCase):
    def setUp(self):
        for entry in os.scandir(UPLOADS):
            shutil.rmtree(entry.path) if entry.is_dir() else os.remove(entry.path)
        server._FAILED_LOGINS.clear()
        # Las pruebas del formato anterior lo analizan; por defecto queda fuera
        model.INCLUDE_LEGACY = True
        self.client = server.app.test_client()

    def put(self, name, content):
        Path(UPLOADS, name).write_bytes(content)

    def upload(self, name, content, headers=VISOR):
        return self.client.post("/upload", headers=headers, data={"file": (io.BytesIO(content), name)})

    def login(self):
        return self.client.post("/login", data={"password": "clave-panel"})

    def status(self, path):
        with self.client.get(path, headers=BROWSER) as response:
            return response.status_code

    def csrf(self):
        page = self.client.get("/sesiones", headers=BROWSER).get_data(as_text=True)
        return re.search(r'name="csrf" value="([^"]+)"', page).group(1)

    def rows(self, table, query=""):
        body = self.client.get(f"/exportar/{table}.csv{query}").get_data(as_text=True)
        return list(csv.DictReader(io.StringIO(body.lstrip("\ufeff"))))


class UploadTest(ServerCase):
    """El contrato con el visor: no puede cambiar sin recompilar el APK."""

    def test_needs_the_api_key(self):
        self.assertEqual(self.upload("0101_5.json", sample("0101_1.json"), headers={}).status_code, 401)
        self.assertEqual(self.upload("0101_5.json", sample("0101_1.json"), {"X-API-Key": "otra"}).status_code, 401)
        self.assertEqual(os.listdir(UPLOADS), [])

    def test_saves_with_the_name_the_headset_sends(self):
        response = self.upload("0101_5.json", sample("0101_1.json"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json(), {"status": "ok", "saved_as": "0101_5.json"})
        self.assertEqual(Path(UPLOADS, "0101_5.json").read_bytes(), sample("0101_1.json"))

    def test_rejects_other_files(self):
        response = self.upload("notas.txt", b"hola")
        self.assertEqual((response.status_code, response.get_json()), (400, {"error": "must be .json or .csv"}))
        response = self.client.post("/upload", headers=VISOR, data={"otro": "1"})
        self.assertEqual((response.status_code, response.get_json()), (400, {"error": "no file field"}))

    def test_name_cannot_leave_the_folder_or_touch_panel_files(self):
        self.assertEqual(self.upload("../../fuera.json", sample("0101_1.json")).get_json()["saved_as"], "fuera.json")
        self.assertEqual(self.upload(".codea2-panel.json", b"{}").get_json()["saved_as"], "codea2-panel.json")
        self.assertEqual(sorted(os.listdir(UPLOADS)), ["codea2-panel.json", "fuera.json"])

    def test_another_session_with_the_same_name_is_kept_apart(self):
        # Dos visores con el mismo PIN y el mismo contador de sesión mandan el mismo nombre
        first, second = sample("0101_1.json"), other_start(sample("0101_1.json"))
        self.upload("0101_5.json", first)
        self.assertEqual(self.upload("0101_5.json", second).get_json()["saved_as"], "0101_5-2.json")
        self.assertEqual(self.upload("0101_5.json", second).get_json()["saved_as"], "0101_5-2.json")
        self.assertEqual(Path(UPLOADS, "0101_5.json").read_bytes(), first)
        self.assertEqual(sorted(os.listdir(UPLOADS)), ["0101_5-2.json", "0101_5.json"])

    def test_reupload_of_the_same_session_keeps_the_previous_version(self):
        first = sample("0101_1.json")
        later = first.replace(b'"blocksGrabbed": 9', b'"blocksGrabbed": 12')
        self.upload("0101_5.json", first)
        self.upload("0101_5.json", first)
        self.assertFalse(Path(UPLOADS, ".versiones").exists())

        self.assertEqual(self.upload("0101_5.json", later).get_json()["saved_as"], "0101_5.json")
        self.assertEqual(Path(UPLOADS, "0101_5.json").read_bytes(), later)
        versions = list(Path(UPLOADS, ".versiones").iterdir())
        self.assertEqual([v.read_bytes() for v in versions], [first])

    def test_health(self):
        response = self.client.get("/health")
        self.assertEqual((response.status_code, response.get_json()), (200, {"status": "alive"}))


class ReadingTest(ServerCase):
    def test_truncated_json_is_recovered_and_flagged(self):
        self.put("prueba_real.json", sample("prueba_real.json"))
        [run] = server.dataset.runs()
        self.assertEqual(run.file.status, "repaired")
        self.assertEqual((run.pin, run.scenarios_completed), ("0103", 4))
        self.assertIn("json_cortado", [issue.code for issue in run.issues])

    def test_file_with_bom_is_read(self):
        self.put("0101_5.json", b"\xef\xbb\xbf" + sample("0101_1.json"))
        [run] = server.dataset.runs()
        self.assertEqual((run.file.status, run.pin, run.issues), ("ok", "0101", ()))

    def test_unreadable_files_are_listed_but_do_not_count(self):
        self.put("0101_5.json", sample("0101_1.json"))
        self.put("roto.json", b"esto no es json")
        self.put("vacio.json", b"")
        self.put("otra_cosa.json", b'{"hola": 1}')
        self.put("viejo.csv", b"a,b\n1,2\n")
        self.assertEqual([r.id for r in server.dataset.runs()], ["0101_5"])
        status = {f.name: f.status for f in server.store.files()}
        self.assertEqual(status, {"0101_5.json": "ok", "roto.json": "invalid", "vacio.json": "invalid",
                                  "otra_cosa.json": "not_run", "viejo.csv": "other"})
        self.login()
        for path in ("/?dif=1", "/exploracion", "/sesiones", "/escenarios", "/datos", "/archivos"):
            self.assertEqual(self.status(path), 200, path)

    def test_session_with_missing_pieces_still_renders(self):
        data = json.loads(sample("0101_1.json"))
        del data["escenario3"]
        data["escenario2"].update(completed=False, endedUtc="")
        data["escenario1"]["attempts"][0]["sequence"] = None
        data["endedUtc"] = ""
        self.put("0101_5.json", json.dumps(data).encode())
        self.login()
        for path in ("/?dif=1", "/exploracion", "/sesiones", "/sesiones/0101_5", "/escenarios", "/datos",
                     "/exportar/todo.xlsx"):
            self.assertEqual(self.status(path), 200, path)
        codes = {issue.code for issue in server.dataset.get("0101_5").issues}
        self.assertTrue({"sin_cierre", "sin_registro", "sin_completar"} <= codes, codes)

    def test_flags_what_does_not_add_up(self):
        # Las dos muestras traen sessionId 5 y 6, pero se llaman _1 y _2
        self.put("0101_1.json", sample("0101_1.json"))
        self.put("copia.json", sample("0101_1.json"))
        codes = {run.id: {issue.code for issue in run.issues} for run in server.dataset.runs()}
        self.assertEqual(codes, {"0101_1": {"nombre_distinto", "duplicada"},
                                 "copia": {"nombre_distinto", "duplicada"}})


class PanelTest(ServerCase):
    def setUp(self):
        super().setUp()
        self.put("0101_5.json", sample("0101_1.json"))
        self.put("0102_6.json", sample("0102_2.json"))

    def test_pages_ask_for_the_password(self):
        response = self.client.get("/sesiones?dif=1", headers=BROWSER)
        self.assertEqual((response.status_code, response.headers["Location"]), (302, "/login?next=/sesiones?dif%3D1"))
        self.assertEqual(self.client.get("/exportar/sesiones.csv").status_code, 401)

    def test_scripts_download_with_basic_auth(self):
        def basic(password):
            return {"Authorization": "Basic " + base64.b64encode(f"codea:{password}".encode()).decode()}

        self.assertEqual(self.client.get("/exportar/sesiones.csv", headers=basic("clave-panel")).status_code, 200)
        # La clave del visor va dentro del APK: sube, pero no lee
        self.assertEqual(self.client.get("/exportar/sesiones.csv", headers=basic("clave-visor")).status_code, 401)

    def test_login(self):
        self.assertEqual(self.client.post("/login", data={"password": "mala"}).status_code, 401)
        self.assertEqual(self.client.get("/", headers=BROWSER).status_code, 302)
        response = self.client.post("/login", data={"password": "clave-panel", "next": "/datos"})
        self.assertEqual((response.status_code, response.headers["Location"]), (302, "/datos"))
        self.assertEqual(self.client.get("/sesiones", headers=BROWSER).status_code, 200)

    def test_login_never_redirects_to_another_site(self):
        for target in ("https://malo.example/", "//malo.example/", "javascript:alert(1)"):
            response = self.client.post("/login", data={"password": "clave-panel", "next": target})
            self.assertEqual(response.headers["Location"], "/", target)
            self.client.post("/logout")

    def test_too_many_wrong_passwords_lock_the_login(self):
        for _ in range(server._MAX_FAILURES):
            self.client.post("/login", data={"password": "mala"})
        self.assertEqual(self.login().status_code, 401)

    def test_every_page_and_download_works(self):
        self.login()
        pages = ["/?dif=1", "/exploracion", "/sesiones", "/sesiones/0101_5", "/escenarios", "/datos", "/archivos", "/archivos/0101_5.json",
                 "/exportar/todo.xlsx", "/exportar/todo-csv.zip", "/exportar/originales.zip",
                 "/?dif=1&desde=2026-09-30&hasta=2026-09-30&pin=01*"]
        pages += [f"/exportar/{t.name}.{fmt}" for t in server.tables.build_tables([]) for fmt in ("csv", "json", "xlsx")]
        for path in pages:
            self.assertEqual(self.status(path), 200, path)
        for path in ("/sesiones/no_existe", "/exportar/no_existe.csv", "/exportar/sesiones.pdf",
                     "/archivos/.codea2-panel.json"):
            self.assertEqual(self.status(path), 404, path)

    def test_excluded_session_leaves_the_analysis(self):
        self.login()
        token = self.csrf()
        response = self.client.post("/sesiones/0102_6/estado", data={"csrf": token, "excluded": "1", "note": "prueba interna"})
        self.assertEqual(response.status_code, 302)

        self.assertEqual([r["file"] for r in self.rows("sesiones")], ["0101_5.json"])
        self.assertEqual({r["file"] for r in self.rows("intentos")}, {"0101_5.json"})
        with_excluded = {r["file"]: (r["excluded"], r["note"]) for r in self.rows("sesiones", "?excluidas=1")}
        self.assertEqual(with_excluded, {"0101_5.json": ("0", ""), "0102_6.json": ("1", "prueba interna")})

        # La nota sigue ahí al volver a incluirla
        self.client.post("/sesiones/0102_6/estado", data={"csrf": token, "excluded": "0"})
        self.assertEqual([(r["excluded"], r["note"]) for r in self.rows("sesiones")], [("0", ""), ("0", "prueba interna")])

    def test_changes_need_the_form_token(self):
        self.login()
        self.assertEqual(self.client.post("/sesiones/0102_6/estado", data={"excluded": "1"}).status_code, 400)
        self.assertEqual(self.client.post("/archivos/subir", data={}).status_code, 400)
        self.assertEqual(len(self.rows("sesiones")), 2)

    def test_manual_upload_from_the_panel(self):
        self.login()
        response = self.client.post("/archivos/subir", data={
            "csrf": self.csrf(), "files": [(io.BytesIO(sample("prueba_real.json")), "0103_4.json"),
                                           (io.BytesIO(b"x"), "foto.png")]})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(sorted(os.listdir(UPLOADS)), ["0101_5.json", "0102_6.json", "0103_4.json"])


class TablesTest(ServerCase):
    def setUp(self):
        super().setUp()
        self.put("0101_5.json", sample("0101_1.json"))
        self.put("0102_6.json", sample("0102_2.json"))
        self.login()

    def test_sessions_table(self):
        first, second = self.rows("sesiones")
        self.assertEqual((first["file"], first["pin"], first["sessionId"], first["difficulty"]), ("0101_5.json", "0101", "5", "1"))
        self.assertEqual((first["startedLocal"], first["date"]), ("2026-09-30 10:15:47", "2026-09-30"))
        self.assertEqual((first["sessionSeconds"], first["completedAll"]), ("300.588", "1"))
        self.assertEqual((first["e1_attempts"], first["e1_totalSeconds"], first["e1_thinkSeconds"]), ("1", "32.946", "28.703"))
        self.assertEqual((first["e2_wrongSelections"], first["e2_deselections"]), ("2", "2"))
        self.assertEqual((second["e3_attempts"], second["e4_wrongPlacements"]), ("2", "4"))

    def test_event_tables_have_one_row_per_event(self):
        counts = {name: len(self.rows(name)) for name in ("escenarios", "intentos", "bloques", "selecciones", "colocaciones")}
        self.assertEqual(counts, {"escenarios": 8, "intentos": 6, "bloques": 32, "selecciones": 10, "colocaciones": 15})
        attempt = self.rows("intentos")[1]
        self.assertEqual((attempt["scenario"], attempt["sequence"], attempt["repetitions"], attempt["itemType"], attempt["solved"]),
                         ("3", "Recoger > Girar al destino > Soltar > Volver", "7", "Caja", "0"))
        # Lo que no aplica queda vacío, no a cero
        self.assertEqual(self.rows("intentos")[0]["repetitions"], "")
        self.assertEqual(self.rows("escenarios")[1]["failedAttempts"], "")

    def test_timeline_is_in_order(self):
        events = [r for r in self.rows("eventos") if r["file"] == "0101_5.json"]
        self.assertEqual((events[0]["event"], events[-1]["event"]), ("inicio_sesion", "fin_sesion"))
        seconds = [float(r["secondsFromSessionStart"]) for r in events]
        self.assertEqual(seconds, sorted(seconds))

    def test_filters_scope_the_downloads(self):
        self.assertEqual(len(self.rows("sesiones", "?dif=2")), 0)
        self.assertEqual([r["pin"] for r in self.rows("sesiones", "?pin=0102")], ["0102"])
        self.assertEqual([r["pin"] for r in self.rows("sesiones", "?pin=01*+-0102")], ["0101"])
        self.assertEqual(len(self.rows("sesiones", "?desde=2026-09-30&hasta=2026-09-30")), 2)
        self.assertEqual(len(self.rows("sesiones", "?desde=2026-10-01")), 0)

    def test_workbook_keeps_types(self):
        book = load_workbook(io.BytesIO(self.client.get("/exportar/todo.xlsx").data))
        self.assertEqual(book.sheetnames, ["analisis", "sesiones", "escenarios", "intentos", "bloques", "selecciones",
                                           "colocaciones", "eventos", "avisos", "diccionario"])
        sheet = book["sesiones"]
        row = dict(zip([c.value for c in sheet[1]], sheet[2]))
        self.assertEqual((row["pin"].value, row["pin"].data_type), ("0101", "s"))
        self.assertEqual((row["sessionId"].value, row["e1_totalSeconds"].value), (5, 32.946))
        self.assertEqual(row["startedLocal"].value.isoformat(), "2026-09-30T10:15:47")

    def test_text_from_the_json_cannot_become_html_or_formulas(self):
        data = json.loads(sample("0101_1.json"))
        data["pin"] = "<b>pin</b>"
        data["escenario2"]["selections"][0]["option"] = '=HYPERLINK("http://malo")'
        data["escenario2"]["selections"][1]["option"] = "<script>alert(1)</script>"
        self.put("hostil.json", json.dumps(data).encode())

        for path in ("/?dif=1", "/exploracion", "/sesiones", "/sesiones/hostil", "/escenarios", "/datos"):
            page = self.client.get(path, headers=BROWSER).get_data(as_text=True)
            self.assertNotIn("<script>alert", page, path)
            self.assertNotIn("<b>pin</b>", page, path)
        self.assertIn("'=HYPERLINK(\"http://malo\")", [r["option"] for r in self.rows("selecciones")])
        book = load_workbook(io.BytesIO(self.client.get("/exportar/selecciones.xlsx").data))
        formulas = [c for row in book["selecciones"].iter_rows() for c in row if c.data_type == "f"]
        self.assertEqual(formulas, [])

    def test_fields_this_server_does_not_know_are_still_exported(self):
        data = json.loads(sample("0101_1.json"))
        data["appVersion"] = "1.4.0"
        data["escenario3"]["repetitionPresses"] = 14
        data["escenario1"]["attempts"][0]["hintsShown"] = 2
        self.put("0101_5.json", json.dumps(data).encode())
        self.assertEqual(self.rows("sesiones")[0]["appVersion"], "1.4.0")
        self.assertEqual(self.rows("escenarios")[2]["repetitionPresses"], "14")
        self.assertEqual(self.rows("intentos")[0]["hintsShown"], "2")

    def test_zips(self):
        tables = zipfile.ZipFile(io.BytesIO(self.client.get("/exportar/todo-csv.zip").data))
        self.assertIn("codea2_sesiones.csv", tables.namelist())
        originals = zipfile.ZipFile(io.BytesIO(self.client.get("/exportar/originales.zip").data))
        self.assertEqual(sorted(originals.namelist()), ["0101_5.json", "0102_6.json"])
        self.assertEqual(originals.read("0101_5.json"), sample("0101_1.json"))



class CurrentFormatTest(ServerCase):
    """El formato del 06/10/2026: visor en la raíz, banderas booleanas, sin contadores que repitan una lista."""

    def setUp(self):
        super().setUp()
        model.INCLUDE_LEGACY = False
        for name in ("0101_5_00000.json", "0102_6_00000.json"):
            self.put(name, (CURRENT / name).read_bytes())
        self.login()

    def test_sessions_are_read_without_warnings(self):
        runs = server.dataset.runs()
        self.assertEqual([(r.pin, r.device_id, r.session_id, r.scenarios_completed) for r in runs],
                         [("0101", "00000", 5, 4), ("0102", "00000", 6, 4)])
        self.assertEqual([r.issues for r in runs], [(), ()])

    def test_tables(self):
        first = self.rows("sesiones")[0]
        self.assertEqual((first["file"], first["deviceId"], first["totalSeconds"]), ("0101_5_00000.json", "00000", "300.588"))
        self.assertEqual((first["e1_blocksConnected"], first["e1_blockResets"], first["e1_errorInvalidUse"]), ("8", "0", "0"))
        # Lo que este formato ya no trae queda vacío; lo que se quitó por repetir una lista, se cuenta
        self.assertEqual((first["e1_blocksReleased"], first["e4_chipsReleased"]), ("", ""))
        self.assertEqual((first["e2_wrongSelections"], first["e4_wrongPlacements"]), ("2", "3"))

        selection = self.rows("selecciones")[0]
        self.assertEqual((selection["selected"], selection["correct"], selection["durationSeconds"]), ("1", "0", "33.848"))
        self.assertEqual(self.rows("colocaciones")[0]["durationSeconds"], "18.83")
        attempt = self.rows("intentos")[0]
        self.assertEqual((attempt["solved"], attempt["attemptDifficulty"]), ("1", ""))
        self.assertNotIn("deseleccion", {r["event"] for r in self.rows("eventos")})

    def test_pages_render(self):
        for path in ("/?dif=1", "/exploracion", "/sesiones", "/sesiones/0101_5_00000", "/escenarios", "/datos",
                     "/exportar/todo.xlsx"):
            self.assertEqual(self.status(path), 200, path)

    def test_both_formats_live_together(self):
        self.put("0103_4.json", sample("prueba_real.json"))
        # La del formato anterior se guarda y se ve, pero no entra en el análisis
        self.assertEqual(len(self.rows("sesiones")), 2)
        self.assertEqual(len(self.rows("sesiones", "?excluidas=1")), 3)
        self.assertEqual(self.status("/sesiones/0103_4"), 200)
        for path in ("/?dif=1", "/exploracion", "/escenarios", "/datos"):
            self.assertEqual(self.status(path), 200, path)

    def test_name_must_carry_the_device(self):
        self.put("0101_5.json", (CURRENT / "0101_5_00000.json").read_bytes())
        codes = {issue.code for issue in server.dataset.get("0101_5").issues}
        self.assertEqual(codes, {"nombre_distinto", "duplicada"})



class MetricsTest(ServerCase):
    """Las métricas del documento del 09/10/2026, con sus propios ejemplos de cálculo."""

    def setUp(self):
        super().setUp()
        model.INCLUDE_LEGACY = False
        for name in ("0101_5_00000.json", "0102_6_00000.json"):
            self.put(name, (CURRENT / name).read_bytes())
        self.login()

    def analysis(self, query=""):
        return {(r["file"], r["scenario"]): r for r in self.rows("analisis", query)}

    def post_sheet(self, run_id, **fields):
        return self.client.post(f"/sesiones/{run_id}/estado", data={"csrf": self.csrf(), **fields})

    def test_home_asks_for_one_difficulty(self):
        response = self.client.get("/", headers=BROWSER)
        self.assertEqual((response.status_code, response.headers["Location"]), (302, "/?dif=1"))
        page = self.client.get("/?dif=1", headers=BROWSER).get_data(as_text=True)
        for code in ("MC1", "MC2", "E1A", "E1B", "E1C", "E2A", "E2B", "E2C", "E3A", "E3B", "E3C", "E4A", "E4B", "E4C"):
            self.assertIn(f">{code}<", page)

    def test_each_scenario(self):
        rows = self.analysis()
        e1, e2 = rows[("0101_5_00000.json", "1")], rows[("0101_5_00000.json", "2")]
        e3, e4 = rows[("0101_5_00000.json", "3")], rows[("0101_5_00000.json", "4")]
        self.assertEqual((e1["G"], e1["N"], e1["K"], e1["F"], e1["errorPct"]), ("1", "1", "1", "0", "0"))
        # 5 selecciones, 2 incorrectas; la primera de enfriamiento falla y las otras dos aciertan
        self.assertEqual((e2["N"], e2["F"], e2["errorPct"]), ("5", "2", "40"))
        self.assertEqual((e2["firstCorrect"], e2["firstObserved"], e2["firstPct"]), ("2", "3", "66.667"))
        # Dos ejecuciones: de 7 con Caja a 3 con Barril es un cambio del plan, no un fallo
        self.assertEqual((e3["N"], e3["K"], e3["unsolved"], e3["F"], e3["errorPct"]), ("2", "2", "1", "", ""))
        self.assertEqual((e3["planChanges"], e3["pairsEvaluable"], e3["pairsExcluded"]), ("1", "1", "0"))
        self.assertEqual((e4["N"], e4["F"], e4["errorPct"]), ("7", "3", "42.857"))
        self.assertEqual((e4["firstCorrect"], e4["firstObserved"], e4["firstPct"]), ("2", "4", "50"))
        self.assertEqual((e1["totalSeconds"], e1["time_excluded"]), ("32.946", ""))

    def test_what_cannot_be_measured_stays_empty(self):
        data = json.loads((CURRENT / "0101_5_00000.json").read_text(encoding="utf-8-sig"))
        data["endedUtc"] = ""
        data["escenario3"].update(completed=False, endedUtc="")
        data["escenario3"]["attempts"] = [dict(a, solved=False, itemType="") for a in data["escenario3"]["attempts"]]
        data["escenario4"] = {"started": False, "completed": False, "totalSeconds": 0.0, "startedUtc": "",
                              "endedUtc": "", "chipsGrabbed": 0, "placements": []}
        self.put("0109_9_00000.json", json.dumps(dict(data, pin="0109", sessionId=9)).encode())
        rows = self.analysis()
        e3, e4 = rows[("0109_9_00000.json", "3")], rows[("0109_9_00000.json", "4")]
        self.assertEqual((e3["G"], e3["G_excluded"], e3["K"], e3["totalSeconds"]), ("", "sin cierre confirmado", "", ""))
        self.assertEqual((e3["planChanges"], e3["pairsEvaluable"], e3["pairsExcluded"]), ("", "0", "1"))
        self.assertEqual((e4["G"], e4["G_excluded"], e4["N"]), ("", "no iniciado", ""))

        # El motivo de cierre de la ficha decide si lo no terminado cuenta como no logrado
        self.post_sheet("0109_9_00000", closure="tiempo")
        self.assertEqual(self.analysis()[("0109_9_00000.json", "3")]["G"], "0")
        self.post_sheet("0109_9_00000", closure="tecnica")
        self.assertEqual(self.analysis()[("0109_9_00000.json", "3")]["G_excluded"], "interrupción técnica")
        self.assertEqual(self.status("/?dif=1"), 200)

    def test_room_that_runs_out_of_time_is_a_final_result(self):
        # Desde el 12/10/2026 cada sala tiene 5 minutos: al agotarlos se pasa a la siguiente
        data = json.loads((CURRENT / "0101_5_00000.json").read_text(encoding="utf-8-sig"))
        data["escenario2"].update(completed=False, timedOut=True)
        data["endedUtc"] = ""
        self.put("0109_9_00000.json", json.dumps(dict(data, pin="0109", sessionId=9)).encode())
        row = self.analysis()[("0109_9_00000.json", "2")]
        self.assertEqual((row["G"], row["timedOut"], row["G_excluded"], row["time_excluded"]),
                         ("0", "1", "", "tiempo agotado"))
        self.assertNotIn("sin_completar", {i.code for i in server.dataset.get("0109_9_00000").issues})
        self.assertEqual({r["file"]: r["e2_timedOut"] for r in self.rows("sesiones")}["0109_9_00000.json"], "1")
        self.assertEqual(self.status("/sesiones/0109_9_00000"), 200)

    def test_session_sheet(self):
        response = self.post_sheet("0101_5_00000", participant=" P-07 ", moment="final", help="verbal",
                                   closure="completada", gameVersion="1.4", p_patrones_1="1", p_patrones_2="0",
                                   p_abstraccion_1="1", p_abstraccion_2="NO", note="tranquilo")
        self.assertEqual(response.status_code, 302)
        row = self.rows("sesiones")[0]
        self.assertEqual((row["participant"], row["moment"], row["help"], row["closure"], row["gameVersion"], row["note"]),
                         ("P-07", "final", "verbal", "completada", "1.4", "tranquilo"))
        # Solo hay porcentaje con los dos criterios observados
        self.assertEqual((row["p_patrones_pct"], row["p_abstraccion_pct"], row["p_algoritmico_pct"]), ("50", "", ""))
        self.assertEqual(self.analysis()[("0101_5_00000.json", "1")]["participant"], "P-07")

        # Tocar otra cosa no borra la ficha, y un valor inventado no se guarda
        self.post_sheet("0101_5_00000", excluded="0", moment="nunca")
        row = self.rows("sesiones")[0]
        self.assertEqual((row["participant"], row["moment"]), ("P-07", ""))
        page = self.client.get("/sesiones/0101_5_00000", headers=BROWSER).get_data(as_text=True)
        self.assertIn('value="P-07"', page)


if __name__ == "__main__":
    unittest.main()
