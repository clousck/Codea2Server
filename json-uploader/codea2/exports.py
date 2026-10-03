"""Descargas: las tablas en CSV, JSON y Excel, y los archivos originales en un ZIP."""

import csv
import io
import json
import math
import os
import zipfile

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .model import as_flag, as_int
from .store import STATE_FILE

# Decimales de los segundos: un float de Unity trae ruido más allá de las milésimas
DECIMALS = 3

# Una celda de texto que empieza así se ejecuta como fórmula al abrir el CSV en Excel
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def _native(value, kind):
    """El valor con su tipo, igual para todos los formatos. None es 'no aplica'."""
    if value is None or value == "":
        return None
    if kind == "flag":
        return 1 if as_flag(value) else 0
    if kind == "int":
        number = as_int(value)
        return str(value) if number is None else number
    if kind == "float":
        try:
            number = float(value)
        except (TypeError, ValueError):
            return str(value)
        # NaN e Infinity no son un dato, y Excel da por roto el libro que los trae
        return round(number, DECIMALS) if math.isfinite(number) else None
    if kind in ("local", "date"):
        return value
    return str(value)


def plain(value, kind):
    """El valor como texto, tal como va en el CSV: punto decimal y banderas 0/1."""
    value = _native(value, kind)
    if value is None:
        return ""
    if kind == "local":
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if kind == "date":
        return value.isoformat()
    if isinstance(value, float):
        text = f"{value:.{DECIMALS}f}".rstrip("0").rstrip(".")
        return "0" if text in ("-0", "") else text
    return str(value)


def to_csv(table):
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(c.name for c in table.columns)
    for row in table.rows:
        cells = []
        for column in table.columns:
            text = plain(row.get(column.name), column.kind)
            if column.kind == "text" and text.startswith(_FORMULA_START):
                text = "'" + text
            cells.append(text)
        writer.writerow(cells)
    # Con BOM, para que Excel lea los acentos; pandas y R lo ignoran
    return ("\ufeff" + buffer.getvalue()).encode("utf-8")


def to_records(table):
    records = []
    for row in table.rows:
        record = {}
        for column in table.columns:
            value = _native(row.get(column.name), column.kind)
            record[column.name] = plain(value, column.kind) if column.kind in ("local", "date") and value else value
        records.append(record)
    return records


def to_json(table):
    return json.dumps(to_records(table), ensure_ascii=False, indent=2).encode("utf-8")


_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1C5CAB")
_FORMATS = {"local": "yyyy-mm-dd hh:mm:ss", "date": "yyyy-mm-dd"}


def to_xlsx(tables):
    """Un libro con una hoja por tabla. Los tipos van nativos: el PIN como texto, las horas como fecha."""
    book = Workbook()
    book.remove(book.active)
    for table in tables:
        sheet = book.create_sheet(table.name)
        widths = []
        for index, column in enumerate(table.columns, start=1):
            cell = sheet.cell(row=1, column=index, value=column.name)
            cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
            cell.alignment = Alignment(vertical="center")
            widths.append(len(column.name))
        for line, row in enumerate(table.rows, start=2):
            for index, column in enumerate(table.columns, start=1):
                value = _native(row.get(column.name), column.kind)
                if value is None:
                    continue
                cell = sheet.cell(row=line, column=index)
                if isinstance(value, str):
                    cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)
                    cell.data_type = "s"  # nunca una fórmula, aunque empiece por "="
                else:
                    cell.value = value
                if column.kind in _FORMATS:
                    cell.number_format = _FORMATS[column.kind]
                widths[index - 1] = max(widths[index - 1], len(plain(value, column.kind)))
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = min(width + 3, 60)
        sheet.freeze_panes = "A2"
        if table.rows:
            sheet.auto_filter.ref = sheet.dimensions
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def csv_zip(tables, prefix="codea2_"):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for table in tables:
            archive.writestr(f"{prefix}{table.name}.csv", to_csv(table))
    return buffer.getvalue()


def originals_zip(store):
    """Todo lo que hay en la carpeta de subidas, tal cual llegó, más el estado del panel."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        names = [f.name for f in store.files()]
        if os.path.exists(store.path(STATE_FILE)):
            names.append(STATE_FILE)
        for name in names:
            archive.write(store.path(name), name)
    return buffer.getvalue()
